"""Build the Paper 1 JOM author-review package without altering frozen evidence.

Default rendering uses the Codex bundled LibreOffice and document renderer. Portable
hosts may specify --render-python and --render-script. Scientific code always runs
with the interpreter used to invoke this script; no system Python is selected.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import build_jeet_submission as common
import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission/jom"
TMP = ROOT / "tmp/jom/render"
RUNTIME = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
RENDERER = (
    Path.home()
    / ".codex/plugins/cache/openai-primary-runtime/documents"
    / "26.1007.11041/skills/documents/render_docx.py"
)
FONT = "Liberation Serif"
TITLE = "Frozen Cross-Dataset Reliability of Healthy-Only PMSM Stator-Fault Detection"
ROMAN = ("I", "II", "III", "IV", "V", "VI")
TABLE_TITLES = (
    "Datasets and fixed health roles",
    "Exploratory KAIST comparison",
    "Frozen external comparison on one physical motor",
)
JOURNAL_ABBREVIATIONS = {
    "IEEE Transactions on Industrial Electronics": "IEEE Trans. Ind. Electron.",
    "IEEE Transactions on Industry Applications": "IEEE Trans. Ind. Appl.",
    "IEEE Transactions on Instrumentation and Measurement": "IEEE Trans. Instrum. Meas.",
    "IEEE Transactions on Power Electronics": "IEEE Trans. Power Electron.",
    "IEEE Transactions on Pattern Analysis and Machine Intelligence": "IEEE Trans. Pattern Anal. Mach. Intell.",
    "Journal of Magnetics": "J. Magn.",
    "Data in Brief": "Data Brief",
    "Advanced Engineering Informatics": "Adv. Eng. Inform.",
    "Mechanical Systems and Signal Processing": "Mech. Syst. Signal Process.",
    "Reliability Engineering & System Safety": "Reliab. Eng. Syst. Saf.",
    "Pattern Recognition": "Pattern Recognit.",
    "IEEE Access": "IEEE Access",
    "IEEE Journal of Emerging and Selected Topics in Power Electronics": "IEEE J. Emerg. Sel. Topics Power Electron.",
    "IEEE Transactions on Industrial Informatics": "IEEE Trans. Ind. Inform.",
    "IEICE Electronics Express": "IEICE Electron. Express",
    "Engineering Applications of Artificial Intelligence": "Eng. Appl. Artif. Intell.",
    "Expert Systems with Applications": "Expert Syst. Appl.",
    "Journal of Process Control": "J. Process Control",
    "Neural Computation": "Neural Comput.",
    "International Journal of Computer Vision": "Int. J. Comput. Vis.",
    "Statistics and Computing": "Stat. Comput.",
    "Journal of Machine Learning Research": "J. Mach. Learn. Res.",
    "Foundations and Trends in Machine Learning": "Found. Trends Mach. Learn.",
    "Ecological Monographs": "Ecol. Monogr.",
    "Ecography": "Ecography",
    "Magnetic Resonance in Medicine": "Magn. Reson. Med.",
    "Technometrics": "Technometrics",
    "Proceedings of Machine Learning Research": "Proc. Mach. Learn. Res.",
}


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def configure(doc: Document, *, supplement: bool = False) -> int:
    common.configure_doc(doc, supplement=supplement)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.85)
    section.top_margin = section.bottom_margin = Inches(0.8)
    for name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Figure Caption",
        "Table Caption",
        "List Bullet",
    ):
        style = doc.styles[name]
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        style.font.name = FONT
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == "Normal" else 12)
        if name == "Title":
            style.font.size = Pt(16)
        if name in {"Heading 1", "Title"}:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.line_spacing = 1.25 if supplement else 2.0
        style.paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Journal of Magnetics author review preparation"
    return int(6.57 * 1440)


def inline(paragraph, text: str, cites: dict[str, int]) -> None:
    common.add_inline_runs(paragraph, text, cites)
    for run in list(paragraph.runs):
        if run.font.name == "Cambria Math" and "_" in run.text:
            parts = re.split(r"_([A-Za-z0-9]+)", run.text)
            run.text = parts[0]
            previous = run._r
            for i, part in enumerate(parts[1:], 1):
                new = paragraph.add_run(part)
                new.font.name = "DejaVu Math TeX Gyre"
                new.font.size = Pt(11)
                new.font.italic = True
                new.font.subscript = bool(i % 2)
                previous.addnext(new._r)
                previous = new._r
        run.font.name = "DejaVu Math TeX Gyre" if run.font.name == "Cambria Math" else FONT
        run.font.size = Pt(11)


def math_run(text: str, *, upright: bool = False):
    run = OxmlElement("m:r")
    props = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p" if upright else "i")
    props.append(sty)
    run.append(props)
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    return run


def math_group(tag: str, nodes):
    group = OxmlElement(f"m:{tag}")
    for node in nodes:
        group.append(node)
    return group


def sub(base, index: str):
    return math_group(
        "sSub",
        [math_group("e", [base]), math_group("sub", [math_run(index, upright=index == "LE")])],
    )


def sup(base, index: str):
    return math_group("sSup", [math_group("e", [base]), math_group("sup", [math_run(index)])])


def tilde_x():
    props = OxmlElement("m:accPr")
    char = OxmlElement("m:chr")
    char.set(qn("m:val"), "̃")
    props.append(char)
    return math_group("acc", [props, math_group("e", [math_run("x")])])


def frac(numerator, denominator):
    return math_group("f", [math_group("num", numerator), math_group("den", denominator)])


def sum_operator(index: str, upper: str, body):
    props = OxmlElement("m:naryPr")
    char = OxmlElement("m:chr")
    char.set(qn("m:val"), "∑")
    props.append(char)
    location = OxmlElement("m:limLoc")
    location.set(qn("m:val"), "undOvr")
    props.append(location)
    return math_group(
        "nary",
        [
            props,
            math_group("sub", [math_run(index)]),
            math_group("sup", [math_run(upper)]),
            math_group("e", body),
        ],
    )


def add_primary_equation(doc: Document, number: int):
    """Author the five frozen equations as structured editable Office Math.

    The legacy JEET linear-LaTeX replacement is intentionally not used: substring
    replacement confused \\left with \\le and lost mathematical structure.
    """
    sm = lambda: sub(math_run("S"), "m")
    ridge = lambda: sup(sm(), "(λ)")
    sigma = lambda: sub(math_run("Σ"), "LE")
    eqs = {
        1: [
            sub(tilde_x(), "m,i"),
            math_run(" = (", upright=True),
            sub(math_run("x"), "m,i"),
            math_run(" − ", upright=True),
            sub(math_run("c"), "m"),
            math_run(") ⊘ ", upright=True),
            sub(math_run("r"), "m"),
        ],
        2: [
            ridge(),
            math_run(" = ", upright=True),
            sm(),
            math_run(" + ", upright=True),
            math_run("λ"),
            frac(
                [math_run("tr(", upright=True), sm(), math_run(")", upright=True)], [math_run("d")]
            ),
            math_run("I"),
        ],
        3: [
            sigma(),
            math_run(" = exp[", upright=True),
            frac([math_run("1")], [math_run("M")]),
            sum_operator(
                "m=1", "M", [math_run("log(", upright=True), ridge(), math_run(")", upright=True)]
            ),
            math_run("]", upright=True),
        ],
        4: [
            math_run("s(x)"),
            math_run(" = ", upright=True),
            sup(tilde_x(), "T"),
            sup(sigma(), "†"),
            tilde_x(),
        ],
        5: [
            sub(math_run("p"), "b"),
            math_run(" = ", upright=True),
            frac(
                [
                    math_run("1 + ", upright=True),
                    sum_operator(
                        "j=1",
                        "n",
                        [
                            math_run("1(", upright=True),
                            sub(math_run("A"), "j"),
                            math_run(" ≥ ", upright=True),
                            sub(math_run("A"), "b"),
                            math_run(")", upright=True),
                        ],
                    ),
                ],
                [math_run("n"), math_run(" + 1", upright=True)],
            ),
        ],
    }
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    math = math_group("oMath", eqs[number])
    paragraph._p.append(math_group("oMathPara", [math]))
    paragraph.add_run(f"     ({number})")


def reference(entry: common.BibEntry) -> str:
    f = entry.fields
    names = []
    for person in f.get("author", "").split(" and "):
        if "," in person:
            last, first = person.split(",", 1)
            names.append(f"{common.initials(first.strip())} {last.strip()}")
        else:
            names.append(person)
    authors = ", ".join(names)
    year = f.get("year", "")
    title = f.get("title", "")
    journal = f.get("journal", f.get("booktitle", f.get("publisher", "")))
    journal = JOURNAL_ABBREVIATIONS.get(journal, journal)
    details = " ".join(x for x in (journal, f.get("volume", "")) if x)
    pages = f.get("pages", "").replace("--", "-")
    if entry.key == "wang2019negativetransfer":
        # Official CVF accepted-manuscript and DOI-registered IEEE pagination differ.
        # The proceedings title/year/DOI unambiguously identify this conference item.
        pages = ""
    if pages:
        details += ", " + pages
    suffix = f" doi:{f['doi']}" if f.get("doi") else f" {f.get('url', '')}"
    return f"{authors}, {title}, {details} ({year}).{suffix}".strip()


def bibliography(doc: Document, cites: dict[str, int]) -> None:
    entries = common.parse_bibtex(ROOT / "references/key_papers.bib")
    extra = OUT / "jom_references.bib"
    if extra.exists():
        entries.update(common.parse_bibtex(extra))
    doc.add_paragraph("References", "Heading 1")
    for key, number in cites.items():
        if key not in entries:
            raise KeyError(f"Unresolved bibliography key {key}")
        p = doc.add_paragraph(f"[{number}] {reference(entries[key])}")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.keep_together = True


def collect_blocks(source: Path):
    lines = source.read_text().splitlines()
    title, i = common.extract_front_title(lines)
    if lines and lines[0] == "---":
        for line in lines[1:i]:
            if line.startswith("title:") and line.partition(":")[2].strip() not in {"", ">-", "|"}:
                title = line.partition(":")[2].strip().strip('"')
    blocks = []
    while i < len(lines):
        s = lines[i].strip()
        i += 1
        if not s or s.startswith(">"):
            continue
        if s == r"\[":
            eq = []
            while i < len(lines) and lines[i].strip() != r"\]":
                eq.append(lines[i].strip())
                i += 1
            i += 1
            blocks.append(("equation", " ".join(eq)))
            continue
        if s.startswith("|"):
            rows, i = common.parse_table(lines, i - 1)
            blocks.append(("table", rows))
            continue
        match = re.match(r"!\[(.*?)\]\((.*?)\)$", s)
        if match:
            blocks.append(("figure", (match[1], (source.parent / match[2]).resolve())))
            continue
        match = re.match(r"^(#{1,4})\s+(.+)", s)
        if match:
            if len(match[1]) > 1:
                blocks.append(("heading", (len(match[1]) - 1, match[2])))
            continue
        text = [s]
        while i < len(lines):
            s = lines[i].strip()
            if not s or s.startswith(("#", "|", "!", ">", "- ")) or s == r"\[":
                break
            text.append(s)
            i += 1
        blocks.append(("bullet" if text[0].startswith("- ") else "paragraph", " ".join(text)))
    return title, blocks


def add_table(doc, rows, width, caption, number, cites):
    common.add_table(
        doc,
        rows,
        width_dxa=width,
        caption=None,
        table_number=number,
        supplement=False,
        cite_state=cites,
    )
    table = doc.tables[-1]
    # Keep a row intact, without imposing a fixed height or keeping the whole table together.
    for row_number, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                if len(table.rows) <= 4:
                    # LibreOffice can lose the last row of a short table at a
                    # page boundary; keep small tables intact across exports.
                    p.paragraph_format.keep_with_next = row_number < len(table.rows) - 1
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(9)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)
    table._tbl.tblPr.append(borders)


def build_document(
    source: Path, output: Path, *, main: bool = False, supplement: bool = False
) -> dict:
    title, blocks = collect_blocks(source)
    doc = Document()
    width = configure(doc, supplement=supplement)
    doc.add_paragraph(title, "Title")
    cites: dict[str, int] = {}
    figures = []
    tables = []
    equations = 0
    figure_captions = {}
    table_captions = {}
    for kind, value in blocks:
        if kind == "paragraph":
            match = re.match(r"Figure (\d+)\.\s+(.+)", value)
            if match:
                figure_captions[int(match[1])] = match[2]
            match = re.match(r"Table ([IVX]+)\.\s+(.+)", value)
            if match:
                table_captions[match[1]] = match[2]
    marked_figures = set()
    for kind, value in blocks:
        if kind == "heading":
            level, heading = value
            if main and heading == "Figure captions":
                continue
            doc.add_paragraph(heading, f"Heading {min(level, 3)}")
        elif kind == "equation":
            equations += 1
            if main:
                add_primary_equation(doc, equations)
            else:
                common.add_o_math(doc, value)
        elif kind == "figure":
            figures.append((figure_captions.get(len(figures) + 1, value[0]), value[1]))
            if not main:
                common.add_figure(doc, value[1], "", len(figures), 6.2)
                cap = doc.paragraphs[-1]
                cap._element.getparent().remove(cap._element)
        elif kind == "table":
            tables.append(value)
            if not main:
                add_table(doc, value, width, None, len(tables), cites)
            else:
                doc.add_paragraph(f"[Table {ROMAN[len(tables) - 1]} near here]")
        else:
            if main and re.match(r"(?:Figure \d+|Table [IVX]+)\.\s", value):
                continue
            paragraph = doc.add_paragraph(style="List Bullet" if kind == "bullet" else None)
            inline(paragraph, value.removeprefix("- ") if kind == "bullet" else value, cites)
            if not main and re.match(r"Table S[\d.]+[a-z]?\.\s", value):
                paragraph.paragraph_format.keep_with_next = True
            if main:
                for number in re.findall(r"\b(?:Fig\.|Figure) (\d+)\b", value):
                    if number not in marked_figures:
                        doc.add_paragraph(f"[Fig. {number} near here]")
                        marked_figures.add(number)
    if cites:
        bibliography(doc, cites)
    if main:
        heading = doc.add_paragraph("Figure captions", "Heading 1")
        heading.paragraph_format.page_break_before = True
        for i, (caption, _) in enumerate(figures, 1):
            doc.add_paragraph(f"Fig. {i}. {caption}")
        for i, rows in enumerate(tables, 1):
            heading = doc.add_paragraph(f"Table {ROMAN[i - 1]}", "Heading 1")
            heading.paragraph_format.page_break_before = True
            doc.add_paragraph(table_captions.get(ROMAN[i - 1], TABLE_TITLES[i - 1]))
            add_table(doc, rows, width, None, i, cites)
        for i, (caption, path) in enumerate(figures, 1):
            heading = doc.add_paragraph(f"Fig. {i}", "Heading 1")
            heading.paragraph_format.page_break_before = True
            # Captions are on the separate caption page required by JOM.
            common.add_figure(doc, OUT / "figures" / f"Fig{i}.png", caption, i, 6.2)
            doc.paragraphs[-1]._element.getparent().remove(doc.paragraphs[-1]._element)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return {
        "figures": len(figures),
        "tables": len(tables),
        "references": len(cites),
        "sha256": digest(output.read_bytes()),
    }


def copy_figures() -> list[dict]:
    _, blocks = collect_blocks(OUT / "manuscript.md")
    figures = [b[1] for b in blocks if b[0] == "figure"]
    captions = {}
    for kind, value in blocks:
        if kind == "paragraph":
            match = re.match(r"Figure (\d+)\.\s+(.+)", value)
            if match:
                captions[int(match[1])] = match[2]
    records = []
    corrected_geometry = ROOT / "tmp/jom/label_corrected_figures"
    subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts/make_external_feature_geometry_figure.py"),
            "--summary",
            str(ROOT / "results/external_feature_drift/feature_diagnostic_summary.csv"),
            "--output-dir",
            str(corrected_geometry),
            "--panel-a-title",
            "(a) Single-feature discrimination",
        ],
        cwd=ROOT,
        check=True,
    )
    dest = OUT / "figures"
    dest.mkdir(exist_ok=True)
    color = dest / "color_sources"
    color.mkdir(exist_ok=True)
    for i, (caption, source) in enumerate(figures, 1):
        caption = captions.get(i, caption)
        publication_source = (
            corrected_geometry / "external_feature_geometry.pdf" if i == 6 else source
        )
        for ext in (".pdf", ".png"):
            shutil.copy2(publication_source.with_suffix(ext), color / f"Fig{i}{ext}")
        with Image.open(publication_source.with_suffix(".png")) as im:
            # Grayscale scientific-plot exports avoid default colour-print requests.
            gray = im.convert("L")
            gray.save(dest / f"Fig{i}.tif", compression="tiff_lzw", dpi=(300, 300))
            gray.save(dest / f"Fig{i}.png", dpi=(300, 300))
            gray.convert("RGB").save(dest / f"Fig{i}.pdf", resolution=300)
        records.append(
            {
                "number": i,
                "source": str(source.relative_to(ROOT)),
                "caption": caption,
                "files": [f"Fig{i}{x}" for x in (".pdf", ".png", ".tif")],
                "publication_exports": "grayscale; original colour vector sources retained separately",
            }
        )
    (dest / "captions.txt").write_text(
        "\n\n".join(f"Fig. {r['number']}. {r['caption']}" for r in records)
    )
    # Original vector artwork and the generating Python scripts remain editable sources.
    sources = OUT / "figure_sources"
    sources.mkdir(exist_ok=True)
    for r in records:
        shutil.copy2(ROOT / r["source"], sources / Path(r["source"]).name)
    for ext in (".pdf", ".png"):
        shutil.copy2(
            corrected_geometry / f"external_feature_geometry{ext}",
            sources / f"JOM_external_feature_geometry{ext}",
        )
    records[-1]["publication_label_correction"] = (
        "Panel (a): Single-feature discrimination; data, axes and selected features unchanged"
    )
    records[-1]["publication_vector_source"] = "figure_sources/JOM_external_feature_geometry.pdf"
    for name in (
        "make_paper1_figures.py",
        "make_external_validation_figures.py",
        "make_external_feature_geometry_figure.py",
    ):
        shutil.copy2(ROOT / "scripts" / name, sources / name)
    return records


RESULT_DIRS = (
    "external_health_audit",
    "external_pmsm_validation",
    "external_pmsm_analysis",
    "external_failure_diagnostics",
    "external_seed_sensitivity",
    "external_feature_drift",
    "sampling_rate_sensitivity",
    "transient_feature_build",
    "transient_feature_build_post_reveal_implicit_time",
    "transient_pmsm_validation_post_reveal_200w",
    "oneclass_baselines",
    "log_ridge_selection",
    "healthy_covariance_v0",
    "manuscript_evidence_validation",
)


def bundle_paths():
    yield ROOT / "pyproject.toml"
    for directory in ("src", "configs", "scripts", "tests", "references"):
        for p in sorted((ROOT / directory).rglob("*")):
            if not p.is_file() or p.suffix not in {".py", ".toml", ".yaml", ".yml", ".bib"}:
                continue
            if any(x in p.parts for x in ("__pycache__", ".pytest_cache", ".ruff_cache")):
                continue
            name = p.name.lower()
            relative_text = p.relative_to(ROOT).as_posix().lower()
            if "/torque/" in relative_text:
                continue
            if any(
                x in name
                for x in ("paper2", "paper3", "paper4", "pmsg", "torque", "thermal", "temperature")
            ):
                continue
            if name in {"test_paper_separation.py"}:
                continue
            if directory == "tests" and name in {
                "test_jeet_submission.py",
                "test_supplementary_material.py",
                "test_conditional.py",
                "test_session_anchor.py",
                "test_operating_context.py",
            }:
                continue
            if directory == "src" and name in {
                "conditional.py",
                "session_anchor.py",
                "operating_context.py",
            }:
                continue
            yield p
    for name in (
        "research_protocol.md",
        "external_validation_protocol.md",
        "fault_reveal_log.md",
        "secondary_transient_validation_protocol.md",
        "secondary_transient_reveal_log.md",
        "data_sources.yaml",
        "reference_metadata_audit.md",
        "jom_submission_requirements_2026-10-09.md",
    ):
        yield ROOT / "docs" / name
    for p in (ROOT / "paper").rglob("*"):
        if p.is_file():
            yield p
    for name in RESULT_DIRS:
        for p in sorted((ROOT / "results" / name).rglob("*")):
            if p.name == "input_manifest.csv":
                continue
            if p.is_file() and p.suffix in {".json", ".csv", ".md", ".gz"}:
                yield p
    for name in (
        "manuscript.md",
        "supplementary.md",
        "cover_letter.md",
        "jom_references.bib",
        "Reproducibility_Guide.md",
        "Current_Environment.txt",
    ):
        yield OUT / name


def portable_data(path: Path) -> bytes:
    data = path.read_bytes()
    # Value/compatibility rows and compressed predictions preserve reveal hashes.
    if path.suffix.lower() in {".csv", ".gz"}:
        return data
    if path.suffix.lower() in {".py", ".md", ".json", ".csv", ".toml", ".yaml", ".bib"}:
        text = data.decode("utf-8")
        if path.parent == OUT and path.suffix == ".md":
            metadata_file = OUT / "author_metadata_REQUIRED.yaml"
            author = yaml.safe_load(metadata_file.read_text()) if metadata_file.exists() else {}
            for key, placeholder in (
                ("full_english_name", "[AUTHOR FULL NAME]"),
                ("email", "[AUTHOR EMAIL]"),
                ("telephone", "[AUTHOR TELEPHONE]"),
                ("telephone_as_provided", "[AUTHOR TELEPHONE]"),
                ("orcid", "[AUTHOR ORCID]"),
                ("postal_address", "[AUTHOR POSTAL ADDRESS]"),
            ):
                if author.get(key):
                    text = text.replace(str(author[key]), placeholder)
        text = text.replace(str(ROOT), ".")
        # Preserve original archived evidence byte-for-byte in the repository; only
        # distributed text copies lose obsolete workstation root coordinates.
        text = text.replace("C:\\\\lkc\\\\phd sci\\\\motortrust", ".")
        text = text.replace("C:\\lkc\\phd sci\\motortrust", ".")
        text = text.replace("C:/lkc/phd sci/motortrust", ".")
        text = text.replace("C:\\\\lkc\\\\phd sci", ".")
        text = text.replace("C:\\lkc\\phd sci", ".")
        text = text.replace("C:/lkc/phd sci", ".")
        text = text.replace("/Users/lkc/miniforge3/envs/motortrust/bin/python", "python")
        text = text.replace("/Users/lkc", "~")
        text = re.sub(r"C:\\Users\\[^\\\s\"]+", "USERPROFILE", text)
        data = text.encode("utf-8")
    return data


def build_zip() -> dict:
    output = OUT / "Paper1_Reproducibility_Code.zip"
    manifest = []
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for path in sorted(set(bundle_paths())):
            if not path.is_file():
                raise FileNotFoundError(path)
            relative = path.relative_to(ROOT).as_posix()
            data = portable_data(path)
            if len(data) > 20_000_000:
                raise ValueError(f"Unexpected large code-bundle file {relative}")
            z.writestr(
                zipfile.ZipInfo(relative, (2026, 10, 9, 0, 0, 0)),
                data,
                compress_type=zipfile.ZIP_DEFLATED,
            )
            manifest.append({"path": relative, "sha256": digest(data), "bytes": len(data)})
        readme_data = (
            b"# Paper 1 reproducibility code\n\n"
            b"See submission/jom/Reproducibility_Guide.md for environment, data access, "
            b"and saved-evidence versus raw-signal reproduction boundaries.\n\n"
            b'Install with python -m pip install -e ".[dev,publication]" then run '
            b"python scripts/audit_jom_evidence.py --output tmp/audit.json.\n\n"
            b"The bundle excludes legacy JEET document-package tests, raw/processed "
            b"signal data, private author contact information and Papers 2-4. "
            b"Original value CSV/gzip files retain reveal hashes. Exported text "
            b"metadata has workstation roots removed.\n"
        )
        z.writestr("README.md", readme_data)
        manifest.append(
            {"path": "README.md", "sha256": digest(readme_data), "bytes": len(readme_data)}
        )
        z.writestr("MANIFEST_SHA256.json", json.dumps(manifest, indent=2))
    with zipfile.ZipFile(output) as z:
        if z.testzip():
            raise ValueError("ZIP CRC failure")
        for entry in manifest:
            if digest(z.read(entry["path"])) != entry["sha256"]:
                raise ValueError("ZIP manifest failure")
        leaks = []
        for name in z.namelist():
            if Path(name).suffix in {".md", ".py", ".json", ".csv", ".yaml", ".toml"}:
                text = z.read(name).decode("utf-8")
                if re.search(r"/Users/[A-Za-z][A-Za-z0-9_-]*/", text):
                    leaks.append(name)
        if leaks:
            raise ValueError(f"Local paths remain in ZIP: {leaks}")
    return {
        "file": output.name,
        "files": len(manifest),
        "bytes": output.stat().st_size,
        "sha256": digest(output.read_bytes()),
        "crc_and_manifest": "pass",
    }


def render(path: Path, render_python: Path, renderer: Path) -> dict:
    dest = TMP / path.stem
    dest.mkdir(parents=True, exist_ok=True)
    for stale in dest.glob("page-*.png"):
        stale.unlink()
    subprocess.run(
        [str(render_python), str(renderer), str(path), "--output_dir", str(dest), "--emit_pdf"],
        check=True,
    )
    pdf = dest / f"{path.stem}.pdf"
    if not pdf.exists():
        raise FileNotFoundError(pdf)
    target = path.with_suffix(".pdf")
    shutil.copy2(pdf, target)
    pages = PdfReader(target).pages
    pngs = list(dest.glob("page-*.png"))
    if len(pngs) != len(pages):
        raise ValueError("Render page count mismatch")
    texts = [p.extract_text() or "" for p in pages]
    return {
        "pages": len(pages),
        "pngs": len(pngs),
        "sha256": digest(target.read_bytes()),
        "visual_review": "requires explicit page-image inspection",
        "minimum_text_characters": min(map(len, texts)),
    }


def legacy_doc(path: Path) -> dict:
    soffice = RUNTIME / "bin/override/soffice"
    with tempfile.TemporaryDirectory(prefix="jom-doc-") as temp:
        subprocess.run(
            [
                str(soffice),
                f"-env:UserInstallation={Path(temp).as_uri()}",
                "--headless",
                "--convert-to",
                "doc:MS Word 97",
                "--outdir",
                str(OUT),
                str(path),
            ],
            check=True,
            capture_output=True,
        )
    target = path.with_suffix(".doc")
    if not target.exists() or target.stat().st_size < 1000:
        raise RuntimeError("Legacy DOC conversion did not produce a valid-size file")
    return {
        "file": target.name,
        "bytes": target.stat().st_size,
        "sha256": digest(target.read_bytes()),
        "native_word_review_required": True,
    }


def build(args) -> dict:
    metadata_path = OUT / "build_metadata.json"
    if metadata_path.exists() and not args.overwrite_docx:
        previous = json.loads(metadata_path.read_text())
        for stem, info in previous.get("documents", {}).items():
            file = OUT / f"{stem}.docx"
            if file.exists() and digest(file.read_bytes()) != info["sha256"]:
                raise ValueError(
                    f"Preserving manually edited Word source: {file.name}. "
                    "Save your edits and update the Markdown source before using --overwrite-docx."
                )
    subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts/validate_manuscript_evidence.py"),
            "--output",
            str(ROOT / "tmp/jom/build_evidence_check.json"),
        ],
        cwd=ROOT,
        check=True,
    )
    subprocess.run(
        [sys.executable, str(ROOT / "scripts/audit_jom_evidence.py")], cwd=ROOT, check=True
    )
    meta = {
        "journal": "Journal of Magnetics",
        "checked_date": "2026-10-09",
        "status": "author review preparation; author inputs and official forms pending",
        "interpreter": sys.executable,
        "documents": {},
        "pdfs": {},
    }
    meta["figures"] = copy_figures()
    sources = (
        ("manuscript.md", "Main_Manuscript_AUTHOR_INPUT_REQUIRED", True, False),
        ("supplementary.md", "Supplementary_Material", False, True),
        ("cover_letter.md", "Cover_Letter_AUTHOR_INPUT_REQUIRED", False, True),
    )
    for source, stem, main, supp in sources:
        output = OUT / f"{stem}.docx"
        meta["documents"][stem] = build_document(OUT / source, output, main=main, supplement=supp)
        if not args.no_render:
            meta["pdfs"][stem] = render(output, args.render_python, args.render_script)
    if not args.no_render:
        meta["legacy_doc"] = legacy_doc(OUT / "Main_Manuscript_AUTHOR_INPUT_REQUIRED.docx")
    meta["zip"] = build_zip()
    (OUT / "build_metadata.json").write_text(json.dumps(meta, indent=2))
    print(json.dumps(meta, indent=2))
    return meta


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-render", action="store_true")
    parser.add_argument(
        "--overwrite-docx",
        action="store_true",
        help="Explicitly replace Word files edited since the last recorded build",
    )
    parser.add_argument("--render-python", type=Path, default=RUNTIME / "python/bin/python3")
    parser.add_argument("--render-script", type=Path, default=RENDERER)
    build(parser.parse_args())


if __name__ == "__main__":
    main()
